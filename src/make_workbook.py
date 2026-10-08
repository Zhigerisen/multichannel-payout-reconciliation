"""Build a client-ready Excel workbook from the engine outputs (headless).

Tabs: Dashboard (KPIs + per-channel margin + money-by-status),
Reconciliation (every order), Exceptions (the breaks), Run Log.

Run:  python src/make_workbook.py   ->  out/reconciliation_demo.xlsx
"""
from __future__ import annotations

import csv
import json
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(__file__), "..", "out")

PANEL = "FF101C33"
EMERALD = "FF34D399"
WHITE = "FFFFFFFF"

H = Font(bold=True, color=WHITE, size=11)
BIG = Font(bold=True, size=20)
ACC = Font(bold=True, color="FF0F7A55", size=20)
LBL = Font(color="FF6B7A90", size=9)
HEAD_FILL = PatternFill("solid", fgColor=PANEL)
MONEY = '#,##0.00'
PCT = '0.00"%"'


def rd(cents) -> float:
    return int(cents) / 100


def read_csv(name: str):
    with open(os.path.join(OUT, name), encoding="utf-8") as f:
        return list(csv.DictReader(f))


def table_sheet(ws, rows, money_cols, headers):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = H
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = MONEY
    ws.freeze_panes = "A2"
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16


def _rows(name):
    return [[r["order_id"], r["channel"], r["status"], rd(r["expected"]),
             rd(r["settled"]), rd(r["delta"]), rd(r["impact"])] for r in read_csv(name)]


HEADERS = ["Order ID", "Channel", "Status", "Expected", "Settled", "Delta", "Impact"]


def build_reconciliation(wb):
    ws = wb.create_sheet("Reconciliation")
    table_sheet(ws, _rows("reconciliation.csv"), [4, 5, 6, 7], HEADERS)
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 3).value != "Matched":
            ws.cell(row, 3).font = Font(bold=True, color="FFB91C1C")


def build_exceptions(wb):
    ws = wb.create_sheet("Exceptions")
    table_sheet(ws, _rows("exceptions.csv"), [4, 5, 6, 7], HEADERS)


def build_runlog(wb):
    ws = wb.create_sheet("Run Log")
    ws.append(["Metric", "Value"])
    for c in (1, 2):
        ws.cell(1, c).font = H
        ws.cell(1, c).fill = HEAD_FILL
    for k, v in read_csv("run_log.csv")[0].items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 28


def kpi(ws, row, col, label, value, acc=False, money=False, pct=False):
    ws.cell(row, col, label.upper()).font = LBL
    vc = ws.cell(row + 1, col, value)
    vc.font = ACC if acc else BIG
    if money:
        vc.number_format = MONEY
    if pct:
        vc.number_format = PCT


def build_dashboard(wb, d):
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "Order & Payout Reconciliation - Dashboard").font = Font(bold=True, size=16)
    ws.cell(2, 1, "Synthetic demonstration dataset - reproducible, fixed seed").font = LBL

    cards = [("GMV", rd(d["gmv"]), False, True, False),
             ("Net payout", rd(d["net_payout"]), False, True, False),
             ("Total fees", rd(d["total_fees"]), False, True, False),
             ("Effective fee", d["effective_fee_pct"], False, False, True),
             ("Orders", d["orders"], False, False, False),
             ("Auto-reconciled", d["reconciled_pct"], True, False, True),
             ("Exceptions", d["exceptions"], False, False, False),
             ("Unreconciled found", rd(d["money_found"]), True, True, False)]
    for i, (lab, val, acc, money, pct) in enumerate(cards):
        kpi(ws, 4 + (i // 4) * 3, 1 + (i % 4) * 2, lab, val, acc, money, pct)

    r = 11
    ws.cell(r, 1, "WHERE THE UNRECONCILED MONEY IS").font = Font(bold=True, size=11)
    ws.cell(r, 4, f"recoverable {rd(d['recoverable']):,.0f} / loss-risk {rd(d['loss_risk']):,.0f}").font = LBL
    r += 1
    for k, v in d["money_by_status"].items():
        ws.cell(r, 1, k)
        ws.cell(r, 2, rd(v)).number_format = MONEY
        r += 1

    r += 1
    start = r
    ws.cell(r, 1, "MARGIN PER CHANNEL").font = Font(bold=True, size=11)
    r += 1
    for c, hcell in enumerate(["Channel", "Revenue", "Fees", "COGS", "Net payout",
                               "Margin %", "Orders"], 1):
        cc = ws.cell(r, c, hcell)
        cc.font = H
        cc.fill = HEAD_FILL
    r += 1
    first = r
    for ch, pc in d["per_channel"].items():
        ws.cell(r, 1, ch.title())
        for col, key in ((2, "revenue"), (3, "fees"), (4, "cogs"), (5, "net_payout")):
            ws.cell(r, col, rd(pc[key])).number_format = MONEY
        ws.cell(r, 6, pc["margin_pct"]).number_format = PCT
        ws.cell(r, 7, pc["orders"])
        r += 1
    last = r - 1

    chart = BarChart()
    chart.title = "Revenue vs Net payout by channel"
    chart.height, chart.width = 7, 14
    chart.add_data(Reference(ws, min_col=2, max_col=5, min_row=first - 1, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    ws.add_chart(chart, f"I{start}")

    for c in range(1, 8):
        ws.column_dimensions[get_column_letter(c)].width = 16


def main():
    with open(os.path.join(OUT, "dashboard.json"), encoding="utf-8") as f:
        d = json.load(f)
    wb = Workbook()
    wb.remove(wb.active)
    build_dashboard(wb, d)
    build_reconciliation(wb)
    build_exceptions(wb)
    build_runlog(wb)
    path = os.path.join(OUT, "reconciliation_demo.xlsx")
    wb.save(path)
    print("[OK] wrote", path, "| tabs:", wb.sheetnames)


if __name__ == "__main__":
    main()
